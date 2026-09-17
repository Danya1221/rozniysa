"""Collect fresh multipart replies, including edits and supported price documents."""
import asyncio
import io
import re
import time
from collections import deque
from contextlib import suppress

from telethon import events, functions, utils
from telethon.errors import BotResponseTimeoutError
from telethon.tl import types

from prices import CLOSED, parse_documents, clean, brand_of, ACCESSORY


class SupplierTimeout(RuntimeError):
    pass


def signature(message):
    buttons = tuple(tuple(b.text for b in row) for row in (message.buttons or []))
    document = getattr(message, "document", None)
    return message.raw_text or "", buttons, getattr(document, "id", None)


def button_label(text):
    return clean(re.sub(r"[^\w\s]", " ", text)).casefold()


def find_button(message, wanted):
    wanted = button_label(wanted)
    for row in message.buttons or []:
        for button in row:
            if button_label(button.text or "") == wanted:
                return button
    return None


def table_text(rows):
    """Use a named retail-price column even if stock/SKU columns follow it."""
    lines = []
    price_column = None
    name_columns = None
    for row in rows:
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        labels = [clean(cell).casefold() for cell in cells]
        price_candidates = [index for index, cell in enumerate(labels)
                            if re.fullmatch(r"(?:цена(?: розничная)?|розница|retail(?: price)?|price|от 1 шт)[., ₽$€а-яa-z]*", cell)]
        names = [index for index, cell in enumerate(labels)
                 if re.search(r"наименован|название|товар|модель|product|title|name|бренд|brand|цвет|color|память|memory|sim|регион|страна|состояние", cell)]
        if price_candidates and names:
            price_column = next((index for index in price_candidates if re.search(r"розни|retail|от 1", labels[index])), price_candidates[0])
            name_columns = names
            continue
        if price_column is not None and price_column < len(cells):
            title = " ".join(cells[index] for index in name_columns if index < len(cells) and cells[index])
            if title and cells[price_column]:
                lines.append(title + " — " + cells[price_column])
            elif title:
                lines.append(title)
        elif any(cells):
            lines.append(" ".join(cell for cell in cells if cell))
        if len(lines) > 30000:
            raise RuntimeError("В таблице больше 30000 строк")
    return "\n".join(lines)


def _username(value):
    if not isinstance(value, str):
        return None
    text = value.strip()
    match = re.fullmatch(r"(?:https?://)?t\.me/([A-Za-z0-9_]+)/?", text, re.I)
    if match:
        text = match.group(1)
    text = text.lstrip("@")
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{3,}", text):
        return text
    return None


def _resolved_entity(result):
    peer = result.peer
    if isinstance(peer, types.PeerUser):
        return next((item for item in result.users if item.id == peer.user_id), None)
    if isinstance(peer, types.PeerChannel):
        return next((item for item in result.chats if item.id == peer.channel_id), None)
    if isinstance(peer, types.PeerChat):
        return next((item for item in result.chats if item.id == peer.chat_id), None)
    return None


def _numeric_value(value):
    if isinstance(value, int):
        return value
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value.strip()):
        return int(value.strip())
    return None


def _entity_allowed(entity, expected):
    if expected == "channel":
        return isinstance(entity, (types.Channel, types.Chat))
    if expected == "user":
        return isinstance(entity, types.User)
    return isinstance(entity, (types.User, types.Channel, types.Chat))


def _matches_numeric(entity, numeric):
    """Match both Bot API marked IDs and raw Telegram object IDs.

    A raw channel id like 6781674751 is positive, while Telethon normally marks
    the same channel as -1006781674751. Treat both forms as the same target.
    """
    if isinstance(entity, types.Channel):
        marked = utils.get_peer_id(entity)
        return numeric in {entity.id, marked}
    if isinstance(entity, types.Chat):
        marked = utils.get_peer_id(entity)
        return numeric in {entity.id, marked, -entity.id}
    if isinstance(entity, types.User):
        return numeric in {entity.id, utils.get_peer_id(entity)}
    return False


async def resolve_input_peer(
    client,
    value,
    *,
    label="Telegram peer",
    dialog_limit=1000,
    expected="any",
):
    """Resolve a configured peer without relying on StringSession entity cache.

    Public usernames are resolved directly and include access_hash. Numeric values
    are also matched against current dialogs by both raw id and Telethon/Bot API
    marked id, so a positive raw channel id is never misread as PeerUser.
    """
    username = _username(value)
    if username:
        try:
            result = await client(functions.contacts.ResolveUsernameRequest(username))
            entity = _resolved_entity(result)
            if entity is None:
                raise RuntimeError("Telegram вернул username без полной entity")
            if not _entity_allowed(entity, expected):
                raise RuntimeError(
                    "это не канал/группа" if expected == "channel" else "неподходящий тип Telegram-чата"
                )
            return utils.get_input_peer(entity)
        except Exception as exc:
            raise RuntimeError(
                f"{label}: не удалось открыть @{username}: {type(exc).__name__}: {exc}"
            ) from exc

    numeric = _numeric_value(value)
    direct_error = None

    # For TARGET_CHANNEL a positive raw channel id must be searched as a channel
    # before Telethon gets a chance to interpret the same positive number as a user.
    if not (expected == "channel" and numeric is not None):
        try:
            direct = await client.get_input_entity(value)
            if expected == "channel" and not isinstance(direct, (types.InputPeerChannel, types.InputPeerChat)):
                raise ValueError("получена entity пользователя вместо канала")
            if expected == "user" and not isinstance(direct, types.InputPeerUser):
                raise ValueError("получена entity канала вместо пользователя")
            return direct
        except Exception as exc:
            direct_error = exc

    try:
        async for dialog in client.iter_dialogs(limit=dialog_limit):
            entity = dialog.entity
            if not _entity_allowed(entity, expected):
                continue
            if numeric is not None and _matches_numeric(entity, numeric):
                return utils.get_input_peer(entity)
            if numeric is None:
                try:
                    if utils.get_peer_id(entity) == utils.get_peer_id(value):
                        return utils.get_input_peer(entity)
                except Exception:
                    pass
    except Exception as dialog_error:
        raise RuntimeError(
            f"{label}: не удалось найти {value!r} в Telegram: "
            f"{type(dialog_error).__name__}: {dialog_error}"
        ) from dialog_error

    # One last direct attempt is useful for negative marked IDs already known by
    # Telethon, but never accept a user for TARGET_CHANNEL.
    if direct_error is None:
        try:
            direct = await client.get_input_entity(value)
            if expected == "channel" and not isinstance(direct, (types.InputPeerChannel, types.InputPeerChat)):
                raise ValueError("получена entity пользователя вместо канала")
            if expected == "user" and not isinstance(direct, types.InputPeerUser):
                raise ValueError("получена entity канала вместо пользователя")
            return direct
        except Exception as exc:
            direct_error = exc

    kind_hint = "канал/группу" if expected == "channel" else "чат"
    raise RuntimeError(
        f"{label}: не удалось получить access_hash для {value!r}. "
        f"Аккаунт должен видеть этот {kind_hint}; поддерживаются @username, -100... и raw channel ID. "
        f"Исходная ошибка: {type(direct_error).__name__}: {direct_error}"
    ) from direct_error


class SupplierReader:
    def __init__(self, client, settings, source):
        self.client = client
        self.settings = settings
        self.source = source
        self.entity = None

    async def resolve(self):
        self.entity = await resolve_input_peer(
            self.client,
            self.source.peer,
            label=self.source.label,
        )

    def is_closed_text(self, text):
        custom = clean(self.source.closed_text).casefold()
        return bool(CLOSED.search(text or "") or (custom and custom in clean(text or "").casefold()))

    async def collect_action(self, action):
        """Snapshot before action, listen before sending, poll as a fallback."""
        s = self.settings
        before = await self.client.get_messages(self.entity, limit=s.history_limit)
        baseline = {m.id: signature(m) for m in before if not m.out}
        max_id = max((m.id for m in before), default=0)
        collected = {}
        changed_at = None

        def accept(message):
            nonlocal changed_at
            if message.out:
                return
            sig = signature(message)
            if message.id in baseline and baseline[message.id] == sig:
                return
            if message.id <= max_id and message.id not in baseline:
                return
            previous = collected.get(message.id)
            if previous is None or signature(previous) != sig:
                collected[message.id] = message
                changed_at = time.monotonic()

        async def on_message(event):
            accept(event.message)

        new_event = events.NewMessage(chats=self.entity, incoming=True)
        edit_event = events.MessageEdited(chats=self.entity, incoming=True)
        self.client.add_event_handler(on_message, new_event)
        self.client.add_event_handler(on_message, edit_event)
        try:
            try:
                await asyncio.wait_for(action(), timeout=s.response_timeout)
            except BotResponseTimeoutError:
                pass
            deadline = time.monotonic() + s.response_timeout
            while time.monotonic() < deadline:
                recent = await self.client.get_messages(self.entity, limit=s.history_limit)
                for message in recent:
                    if message.id > max_id or message.id in baseline:
                        accept(message)
                if collected and changed_at is not None and time.monotonic() - changed_at >= s.quiet_seconds:
                    # Fetch the entire fresh range, including replies beyond the
                    # history window if event delivery lagged behind the supplier.
                    recent = await self.client.get_messages(self.entity, min_id=max_id, limit=None)
                    for message in recent:
                        accept(message)
                    if time.monotonic() - changed_at < s.quiet_seconds:
                        continue
                    return [collected[k] for k in sorted(collected)]
                await asyncio.sleep(s.action_delay)
            if collected:
                raise SupplierTimeout("Ответ поставщика не завершён: увеличь RESPONSE_TIMEOUT")
            raise SupplierTimeout("Поставщик не прислал нового ответа; старое меню не будет опубликовано")
        finally:
            self.client.remove_event_handler(on_message, new_event)
            self.client.remove_event_handler(on_message, edit_event)

    async def messages(self):
        if self.entity is None:
            await self.resolve()
        if self.source.mode == "feed":
            messages = [m for m in await self.client.get_messages(
                self.entity, limit=self.settings.history_limit)]
            if not messages:
                raise SupplierTimeout("Лента поставщика пуста")
            latest = next((m for m in messages if (m.raw_text or "").strip() or m.document), None)
            if latest and self.is_closed_text(latest.raw_text):
                return [latest]
            # Static supplier channels may contain more than 300 live price posts.
            # Unlike the action snapshot, this read must cover the full configured feed.
            if len(messages) >= self.settings.history_limit:
                messages = list(await self.client.get_messages(
                    self.entity, limit=self.settings.feed_history_limit or None))
            return list(reversed(messages))
        if not self.source.request:
            raise ValueError("Для SOURCE_MODE=bot нужен REQUEST_TEXT")
        messages = await self.collect_action(
            lambda: self.client.send_message(self.entity, self.source.request, parse_mode=None))
        for wanted in self.source.buttons:
            if any(self.is_closed_text(m.raw_text) for m in messages):
                return messages
            button = next((b for m in reversed(messages) if (b := find_button(m, wanted))), None)
            if button is None:
                labels = [b.text for m in messages for row in (m.buttons or []) for b in row]
                raise RuntimeError(f"Нет кнопки «{wanted}». Доступны: {', '.join(labels)}")
            if not isinstance(button.button, (types.KeyboardButton, types.KeyboardButtonCallback)):
                raise RuntimeError(f"Кнопка «{wanted}» не является текстовой или callback-кнопкой")
            messages = await self.collect_action(button.click)
        return messages

    async def expand_catalog(self, messages):
        """Read linked price posts and read-only category/pagination buttons."""
        queue = deque(messages)
        result = []
        visited = set()
        resolved = {}
        while queue:
            message = queue.popleft()
            result.append(message)
            for row in message.buttons or []:
                for button in row:
                    raw = button.button
                    label = button_label(button.text)
                    url = getattr(raw, "url", "") or ""
                    link = re.fullmatch(r"https?://(?:t|telegram)\.me/(?:(c)/(\d+)|([A-Za-z0-9_]+))/(\d+)(?:\?[^#]*)?", url)
                    navigation = (brand_of(label) or ACCESSORY.search(label)
                                  or re.fullmatch(r"(?:прайс|каталог|весь прайс|актуальный прайс|далее|следующая|впер[её]д|next|\d+)", label)
                                  or (not label and button.text.strip() in {"→", "➡", "➡️", "▶", "▶️", ">", "»"}))
                    if re.search(r"купить|заказ|корзин|оплат|удал|брон|резерв|buy|order|cart|pay|reserve", label):
                        continue
                    callback = (isinstance(raw, types.KeyboardButtonCallback) and navigation
                                and self.source.mode == "bot" and len(label) <= 60
                                and not re.search(r"\d{4,}|[₽$€]", button.text))
                    if not link and not callback:
                        continue
                    pagination = bool(re.fullmatch(r"(?:далее|следующая|впер[её]д|next|\d+)", label) or not label)
                    key = ("url", url) if link else ("callback", raw.data, signature(message) if pagination else None)
                    if key in visited:
                        continue
                    if len(visited) >= self.settings.catalog_pages:
                        raise RuntimeError("Каталог больше MAX_CATALOG_PAGES; увеличь лимит. Неполный прайс не опубликован")
                    visited.add(key)
                    if link:
                        target = int("-100" + link[2]) if link[1] else "@" + link[3]
                        if target not in resolved:
                            resolved[target] = await resolve_input_peer(self.client, target, label="Раздел прайса", expected="channel")
                        linked = await self.client.get_messages(resolved[target], ids=int(link[4]))
                        if not linked or isinstance(linked, types.MessageEmpty):
                            raise RuntimeError("Поставщик удалил раздел каталога: " + button.text)
                        queue.append(linked)
                    else:
                        queue.extend(await self.collect_action(button.click))
                    await asyncio.sleep(self.settings.action_delay)
        return result

    async def document_text(self, message):
        document = getattr(message, "document", None)
        if not document:
            return ""
        if document.size > 5 * 1024 * 1024:
            raise RuntimeError("Файл прайса больше 5 МБ")
        name = next((a.file_name for a in document.attributes if hasattr(a, "file_name")), "").lower()
        if not name.endswith((".txt", ".csv", ".xlsx")):
            raise RuntimeError("Поддерживаются текстовые прайсы, TXT, CSV и XLSX; получен другой файл")
        payload = await self.client.download_media(message, file=bytes)
        if not payload:
            raise RuntimeError("Не удалось скачать файл прайса")
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            import zipfile
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise RuntimeError("Распакованный XLSX превышает 50 МБ")
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                return "\n".join(table_text(sheet.iter_rows(values_only=True)) for sheet in workbook)
            finally:
                workbook.close()
        text = None
        for encoding in ("utf-8-sig", "cp1251"):
            with suppress(UnicodeDecodeError):
                text = payload.decode(encoding)
                break
        if text is None:
            raise RuntimeError("Не удалось прочитать кодировку TXT/CSV")
        if name.endswith(".csv"):
            import csv
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
                rows = csv.reader(io.StringIO(text), dialect)
                text = table_text(rows)
            except csv.Error:
                pass
        return text

    async def fetch(self):
        messages = await self.expand_catalog(await self.messages())
        documents = []
        for message in messages:
            if message.raw_text:
                documents.append(message.raw_text)
            if message.document:
                documents.append(await self.document_text(message))
        result = parse_documents(documents, self.settings.currency)
        if not result.items and any(self.is_closed_text(text) for text in documents):
            result.closed = True
        if not result.items and not result.closed and not (result.unavailable and not result.rejected):
            raise RuntimeError("В ответе нет распознанных товаров. Прайс в канале сохранён")
        return result
